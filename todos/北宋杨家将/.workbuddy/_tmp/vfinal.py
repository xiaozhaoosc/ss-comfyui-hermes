# -*- coding: utf-8 -*-
import os, re, traceback
BASE = r"D:\BaiduNetdiskDownload\北宋杨家将\动画改编方案"
LOG = r"D:\BaiduNetdiskDownload\北宋杨家将\.workbuddy\_tmp\qc\vfinal.log"
L = []
try:
    for f in ["杨家将_竖屏版规格书.html", "杨家将_剧本质检报告.html", "杨家将_剧本与分镜全本.html"]:
        p = os.path.join(BASE, f)
        s = open(p, "r", encoding="utf-8").read()
        L.append("== %s  (%.1f KB)" % (f, len(s.encode("utf-8")) / 1024))
        toks = re.findall(r"@@[A-Z0-9]+@@", s)
        L.append("   残留 token = %s" % (set(toks) if toks else "无"))
        for t in ("html", "head", "body", "div", "table", "tbody", "thead", "section", "main", "footer", "ol", "ul"):
            o = len(re.findall(r"<%s[ >]" % t, s)); c = len(re.findall(r"</%s>" % t, s))
            if o != c:
                L.append("   *** <%s> %d/%d 不匹配 ***" % (t, o, c))
        L.append("   <tr> %d 行  <table> %d 个" % (len(re.findall(r"<tr[ >]", s)), len(re.findall(r"<table[ >]", s))))
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(BASE, "杨家将_Flow片段切分表.xlsx"))
    L.append("")
    L.append("XLSX sheets: " + " / ".join(wb.sheetnames))
    for n in wb.sheetnames:
        L.append("  [%s] rows=%d cols=%d" % (n, wb[n].max_row, wb[n].max_column))
    ws = wb["全片切分表812段"]
    L.append("  首行样本: " + " | ".join(str(ws.cell(2, c).value)[:26] for c in range(1, 11)))
    durs = {}
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 6).value
        durs[d] = durs.get(d, 0) + 1
    L.append("  段长分布: %s" % durs)
except Exception:
    L.append("FATAL\n" + traceback.format_exc())
open(LOG, "w", encoding="utf-8").write("\n".join(L))
