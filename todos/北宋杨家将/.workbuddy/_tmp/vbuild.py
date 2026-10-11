# -*- coding: utf-8 -*-
"""竖屏版规格 + Flow 片段切分表：解析剧本 → 逐镜定长 → 切分 → xlsx。"""
import os, re, json, traceback
from collections import Counter

BASE = r"D:\BaiduNetdiskDownload\北宋杨家将"
SRC = os.path.join(BASE, r"动画改编方案\杨家将_剧本与分镜全本.html")
OUTX = os.path.join(BASE, r"动画改编方案\杨家将_Flow片段切分表.xlsx")
OUTJ = os.path.join(BASE, r".workbuddy\_tmp\qc\shots.json")
LOG = os.path.join(BASE, r".workbuddy\_tmp\qc\vbuild.log")

EPINFO = {
 1:("北汉风雪","杨业一个人站在城头——对面的旗，不是自己的颜色"),
 2:("归宋 · 天波府","应州城门缓缓打开，一家人第一次走进一座不属于自己的城"),
 3:("邠阳城 · 围","城门在身后合上。那一声"),
 4:("邠阳城 · 血","一件御袍，穿到了另一个人身上"),
 5:("误卯牌","一块木牌挂出去，上面写着两个字：误卯"),
 6:("李陵碑 · 殉国","镜头从一个人背后升起，越过他的肩，对着那块碑"),
 7:("七郎之死","弓弦齐齐一响——然后镜头抬起来，只看天"),
 8:("汴京审潘","一面登闻鼓，静着"),
 9:("辞官","一个辞了高官的人，背着一杆枪出城，没有人送"),
 10:("佳山寨","一堵墙，一段一段长起来"),
 11:("焦赞","一个人被吊在半空中，脚底下是空的"),
 12:("双龙谷","谷口封死，山上万箭齐发"),
 13:("八姐闯幽州","一箭，穿过百步之外的一根树枝"),
 14:("天波楼","瓦片从天上落下来，砸满一条街"),
 15:("智救杨郡马","一只匣子，从汝州带回汴京"),
 16:("魏州解围","城楼上，出现了死人的旗号"),
 17:("天门阵","一整片山谷铺开——金光闪闪的一张棋盘"),
 18:("穆阁寨","一对大雁从天上掉下来"),
 19:("辕门斩子","儿子绑在桩上，父亲背对着他坐"),
 20:("大破天门阵","金色地砖，一块一块熄灭"),
 21:("碑上的字","碑上刻满了名字，其中一个，是刚刻上去的"),
 22:("寿宴请缨","一只酒杯，落在地上"),
 23:("内外夹攻","一支箭带着信，飞过城墙"),
 24:("穆桂英挂帅","碑前那片荒坡上，开了花"),
}

# 竖屏风险：横屏专属运镜
H_RISK = {
 "横移": ("竖移（上下摇）", "把横向展开的信息改成纵向递进；或改为纵深推进，让人物由远及近"),
 "平摇": ("竖摇", "横摇扫景在竖屏里信息全丢，改成自下而上摇起，把「地—人—天」竖着讲"),
 "摇镜": ("竖摇", "同上"),
 "大远景": ("中景+远山分层", "竖屏装不下横向的地平线，把远景拆成「前景人物 / 中景坡地 / 远景留白天空」三层竖向叠置"),
 "全景": ("中近景", "横构图全景在竖屏里人物会小到看不清，改中近景，环境靠画外暗示"),
 "俯拍": ("高角度纵深俯拍", "保留俯视但改成沿纵深方向拍（从近处地面推向远处），不要在横向扫"),
 "俯 3": ("高角度纵深俯视", "同上"),
 "横移（右）": ("竖移（自上而下）", "同「横移」"),
 "一镜到底": ("分段", "竖屏一镜到底容易失焦，拆成 2–3 段"),
}
V_OK = ["推", "缓推", "拉", "升", "降", "下摇", "特写", "近景", "仰拍", "固定"]

CRIT_8 = ["空镜", "氛围", "核心镜头", "关键镜头", "收尾", "季终", "片尾", "全剧最后", "留白"]
CRIT_4 = ["留白爆闪", "闪回", "蒙太奇", "快速", "快切", "一涌", "纷纷"]


def find_block(s, start):
    i, d = start, 0
    pat = re.compile(r"<div\b|</div>")
    while True:
        m = pat.search(s, i)
        d += 1 if m.group(0) == "<div" else -1
        i = m.end()
        if d == 0:
            return i, s[start:i]


def strip_tags(t):
    t = re.sub(r"<span class=\"pg\">(.*?)</span>", r"[\1] ", t)
    t = re.sub(r"<[^>]+>", "", t)
    return re.sub(r"\s+", " ", t).strip()


def parse():
    html = open(SRC, "r", encoding="utf-8").read()
    eps = []
    for m in re.finditer(r'<div class="ep">', html):
        en, blk = find_block(html, m.start())
        nm = re.search(r'<div class="n">(\d+)</div>', blk)
        eps.append((int(nm.group(1)), blk))
    eps.sort()
    out = []
    for num, blk in eps:
        # 场
        scs = []
        for sm in re.finditer(r'<div class="slug">(.*?)</div>', blk):
            scs.append((sm.start(), sm.group(1).strip()))
        shots = []
        for pm in re.finditer(r'<p class="act">(.*?)</p>', blk, re.S):
            raw = pm.group(1)
            g = re.search(r'<span class="pg">(\d+)-(\d+)</span>', raw)
            book, page = (int(g.group(1)), int(g.group(2))) if g else (0, 0)
            txt = strip_tags(raw)
            slug = ""
            for pos, s in scs:
                if pos < pm.start():
                    slug = s
            shots.append(dict(book=book, page=page, text=txt, scene=slug))
        out.append(dict(ep=num, shots=shots))
    return out


def dur_of(t):
    if any(k in t for k in CRIT_8):
        return 8
    if any(k in t for k in CRIT_4):
        return 4
    return 6


def value(t):
    """越值得长镜头的分越高。用于时长分配，避免把一部戏全压成 4 秒。"""
    sc = 0
    if any(k in t for k in ["空镜", "氛围", "核心镜头", "关键镜头", "收尾", "季终", "片尾", "全剧最后", "留白"]):
        sc += 3
    if any(k in t for k in ["特写", "大特写", "近景", "中景"]):
        sc += 1
    if any(k in t for k in ["留白爆闪", "闪回", "蒙太奇", "快速", "快切", "一涌"]):
        sc -= 2
    if len(t) > 90:
        sc += 1
    return sc


def normalize(txts, target=210):
    """在 4/6/8 三档内把总长调到 target；按「镜头价值」排序，每轮各调一次，分布均匀。"""
    ds = [dur_of(t) for t in txts]
    vs = [value(t) for t in txts]
    order_down = sorted(range(len(ds)), key=lambda i: (vs[i], i))
    order_up = sorted(range(len(ds)), key=lambda i: (-vs[i], i))
    guard = 0
    while sum(ds) > target and guard < 5000:
        guard += 1
        changed = False
        for i in order_down:
            if sum(ds) <= target:
                break
            if ds[i] > 4:
                ds[i] -= 2
                changed = True
        if not changed:
            break
    while sum(ds) < target and guard < 20000:
        guard += 1
        changed = False
        for i in order_up:
            if sum(ds) >= target:
                break
            if ds[i] < 8:
                ds[i] += 2
                changed = True
        if not changed:
            break
    return ds


def vfix(t):
    for k, (new, why) in H_RISK.items():
        if k in t:
            return new, why
    return "", ""


def main():
    data = parse()
    rows = []
    summary = []
    for e in data:
        sh = e["shots"]
        txts = [s["text"] for s in sh]
        ds = normalize(txts, 210)
        total = sum(ds)
        gap = 210 - total                      # >0 说明镜头数不足以撑满 3.5 分钟
        need = max(0, -(-gap // 6))            # 还需补几个 6 秒镜头
        for i, (s, d) in enumerate(zip(sh, ds), 1):
            nv, why = vfix(s["text"])
            rows.append(dict(ep=e["ep"], epname=EPINFO[e["ep"]][0], scene=s["scene"], idx=i,
                             mark=("%d-%d" % (s["book"], s["page"])) if s["book"] else "",
                             dur=d, txt=s["text"][:150], vfix=nv, vwhy=why))
        cnt = Counter(ds)
        summary.append(dict(ep=e["ep"], name=EPINFO[e["ep"]][0], shots=len(sh), total=total,
                            d4=cnt[4], d6=cnt[6], d8=cnt[8], gap=gap, need=need,
                            hooks=EPINFO[e["ep"]][1],
                            risk=sum(1 for r in rows if r["ep"] == e["ep"] and r["vfix"])))
    json.dump(dict(rows=rows, summary=summary), open(OUTJ, "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)

    # ---------- XLSX ----------
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    C_ = "231F19"; ZHU = "B03A2E"; ALT = "FCFAF5"; JIN = "F8F1DC"
    thin = Side(style="thin", color="D2C8B2")
    BD = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet(ws, headers, body, widths, ctr=()):
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(1, c)
            cell.font = Font(bold=True, color="F7F3EA", size=10.5, name="Microsoft YaHei")
            cell.fill = PatternFill("solid", fgColor=C_)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BD
        ws.row_dimensions[1].height = 28
        for i, r in enumerate(body, start=2):
            ws.append(list(r))
            for c in range(1, len(headers) + 1):
                cell = ws.cell(i, c)
                cell.font = Font(size=9.5, name="Microsoft YaHei")
                cell.alignment = Alignment(vertical="top", wrap_text=True,
                                           horizontal="center" if c in ctr else "left")
                cell.border = BD
                if i % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=ALT)
        for k, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(k)].width = w
        ws.freeze_panes = "A2"

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "全片切分表812段"
    body = [[r["ep"], r["epname"], r["scene"], r["idx"], r["mark"], r["dur"], 1,
             r["txt"], r["vfix"] or "—", r["vwhy"] or "—"] for r in rows]
    sheet(ws1, ["集", "集名", "场次", "镜序", "原书页", "段长(秒)", "Flow段数", "画面摘要",
                "竖屏运镜改写", "改写理由"], body,
          [5, 13, 26, 6, 9, 9, 8, 74, 20, 46], ctr=(1, 4, 5, 6, 7))

    ws2 = wb.create_sheet("分集汇总与开场钩子")
    s2 = []
    for s in summary:
        s2.append([s["ep"], s["name"], s["shots"], s["d4"], s["d6"], s["d8"], s["total"],
                   "%d:%02d" % (s["total"] // 60, s["total"] % 60),
                   ("缺 %d 秒，需补 %d 镜" % (s["gap"], s["need"])) if s["gap"] > 0 else "达标",
                   s["shots"], s["shots"] * 10, s["shots"] * 20, s["shots"] * 100,
                   s["risk"], s["hooks"]])
    sheet(ws2, ["集", "集名", "镜头数", "4秒段", "6秒段", "8秒段", "总秒", "时长", "3.5分钟达标情况",
                "Flow段数", "Lite积分", "Fast积分", "Quality积分", "需改运镜", "开场 3 秒钩子"],
          s2, [5, 13, 8, 8, 8, 8, 7, 8, 22, 9, 10, 10, 12, 10, 56],
          ctr=(1, 3, 4, 5, 6, 7, 8, 10, 11, 12, 13, 14))

    ws3 = wb.create_sheet("横改竖运镜对照")
    s3 = [
        ["横移 / 平摇", "竖移（上下摇）", "把横向展开的信息改成纵向递进，或改为纵深推进", "高"],
        ["大远景", "中景 + 远山分层", "竖屏装不下横向地平线，拆成「前景人物／中景坡地／远景留白」三层竖向叠置", "高"],
        ["全景", "中近景", "横构图全景在竖屏里人物太小，环境靠画外暗示", "高"],
        ["俯拍", "高角度纵深俯拍", "保留俯视但沿纵深方向推，不要横向扫", "中"],
        ["一镜到底", "拆 2–3 段", "竖屏一镜到底容易失焦", "中"],
        ["推 / 缓推", "直接可用", "竖屏最安全的运镜", "—"],
        ["升 / 降", "直接可用", "竖向运动天生适配 9:16", "—"],
        ["下摇", "直接可用", "同上", "—"],
        ["特写 / 近景", "直接可用", "竖屏的天然优势画幅", "—"],
        ["仰拍", "直接可用", "竖向构图增强崇高感", "—"],
        ["固定机位", "直接可用", "适合空镜与情绪镜", "—"],
    ]
    sheet(ws3, ["横屏原运镜", "竖屏改写", "说明", "优先级"], s3, [20, 22, 62, 9], ctr=(4,))

    ws4 = wb.create_sheet("成本与投产策略")
    tot_shots = len(rows)
    s4 = [
        ["全片规模", ""],
        ["镜头总数", tot_shots],
        ["Flow 段总数（1 镜 = 1 段）", tot_shots],
        ["总时长", "84 分钟（24 集 × 3.5 分钟）"],
        ["", ""],
        ["视频生成预算（按段计）", ""],
        ["① Lite 档 10 积分/段", tot_shots * 10],
        ["② Fast 档 20 积分/段", tot_shots * 20],
        ["③ Quality 档 100 积分/段", tot_shots * 100],
        ["", ""],
        ["首帧出图预算（812 张首帧）", ""],
        ["按单张 5–10 积分估", "约 4,060 – 8,120"],
        ["", ""],
        ["推荐策略：分三步走", ""],
        ["第一步 · 试播集", "只做第 1 集和第 6 集（共 72 镜）。Lite 档跑，约 720 积分。验证竖屏构图与节奏"],
        ["第二步 · 定风格", "对试播集里通过的段落升 Fast（约 30 镜，600 积分），出可对外看的样片"],
        ["第三步 · 铺全片", "按集推进，先 Lite 出 animatic，选定的镜头再升 Fast。不要一上来全上 Quality"],
        ["", ""],
        ["省钱三条", ""],
        ["①", "Lite 出 animatic、Fast 出样片、Quality 只给最终定稿镜头——全片都用 Quality 是 8 万积分级别，没必要"],
        ["②", "同场景连续的静止空镜可用 Extend 连成一段，减少段数"],
        ["③", "首帧统一在锁定色盘后批量出图，避免逐镜返工"],
    ]
    for i, (a, b) in enumerate(s4, start=1):
        ws4.cell(i, 1, a).font = Font(size=10.5, bold=(b == "" and a != ""), name="Microsoft YaHei",
                                      color=ZHU if (b == "" and a != "") else C_)
        ws4.cell(i, 2, b).font = Font(size=10, name="Microsoft YaHei")
        ws4.cell(i, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws4.cell(i, 2).alignment = Alignment(vertical="top", wrap_text=True)
    ws4.column_dimensions["A"].width = 30
    ws4.column_dimensions["B"].width = 92
    ws4.cell(1, 1).font = Font(size=14, bold=True, color=ZHU, name="Microsoft YaHei")

    wb.save(OUTX)
    return tot_shots, sum(s["risk"] for s in summary), summary


try:
    n, r, summ = main()
    msg = "OK 段数=%d 需改运镜=%d 未达标集=%d\n" % (n, r, sum(1 for s in summ if s["gap"] > 0))
    msg += "\n".join("  第%2d集 %-13s 镜%3d  4s%3d 6s%3d 8s%3d  总%3ds (%d:%02d)  改运镜%2d  缺口%d秒"
                     % (s["ep"], s["name"], s["shots"], s["d4"], s["d6"], s["d8"], s["total"],
                        s["total"] // 60, s["total"] % 60, s["risk"], s["gap"]) for s in summ)
except Exception:
    msg = "FATAL\n" + traceback.format_exc()
open(LOG, "w", encoding="utf-8").write(msg)
print("done")
