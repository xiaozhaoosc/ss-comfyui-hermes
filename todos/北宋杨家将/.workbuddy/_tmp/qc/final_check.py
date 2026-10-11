# -*- coding: utf-8 -*-
import os, re, traceback
BASE = r"D:\BaiduNetdiskDownload\北宋杨家将\动画改编方案"
LOG = r"D:\BaiduNetdiskDownload\北宋杨家将\.workbuddy\_tmp\qc\final_check.log"
L = []
def chk(path, name):
    s = open(path, "r", encoding="utf-8").read()
    L.append("== %s  (%.1f KB)" % (name, len(s.encode("utf-8")) / 1024))
    for t in ("html", "head", "body", "div", "table", "tbody", "thead", "section", "main", "footer", "p"):
        o = len(re.findall(r"<%s[ >]" % t, s)); c = len(re.findall(r"</%s>" % t, s))
        flag = "OK" if o == c else "*** 不匹配 ***"
        if flag != "OK":
            L.append("   <%s> %d / %d %s" % (t, o, c, flag))
    L.append("   div=%d/%d  集块=%d  镜头=%d  页标记=%d" %
             (len(re.findall(r"<div[ >]", s)), len(re.findall(r"</div>", s)),
              len(re.findall(r'<div class="ep">', s)), s.count('<p class="act">'),
              len(re.findall(r'class="pg"', s))))
    return s

try:
    a = chk(os.path.join(BASE, "杨家将_剧本与分镜全本.html"), "剧本与分镜全本")
    b = chk(os.path.join(BASE, "杨家将_剧本质检报告.html"), "剧本质检报告")
    # 残留检查
    for kw in ("631 页逐页落镜", "818 个 △", "第 8 集）、<b>桂英挂帅", "余太君（"):
        L.append("  残留「%s」= %d" % (kw, a.count(kw)))
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(BASE, "杨家将_逐页分镜总表.xlsx"))
    ws = wb["分集页数分配"]
    L.append("   XLSX 分集表 rows=%d" % ws.max_row)
    for r in range(2, ws.max_row + 1):
        L.append("     %s 册%s %s 页数%s" % (ws.cell(r,1).value, ws.cell(r,4).value,
                                            ws.cell(r,5).value, ws.cell(r,6).value))
except Exception:
    L.append("FATAL\n" + traceback.format_exc())
open(LOG, "w", encoding="utf-8").write("\n".join(L))
