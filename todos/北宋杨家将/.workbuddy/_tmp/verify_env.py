import os, traceback

LOG = r"D:\BaiduNetdiskDownload\北宋杨家将\.workbuddy\_tmp\verify.log"
lines = []

# 1) pymupdf
try:
    import pymupdf
    lines.append("pymupdf OK: " + pymupdf.__version__)
    d = pymupdf.open(r"D:\BaiduNetdiskDownload\北宋杨家将\漫画\01杨令公归宋.pdf")
    lines.append("open OK, pages = %d" % d.page_count)
    pix = d[0].get_pixmap(matrix=pymupdf.Matrix(0.3, 0.3), colorspace=pymupdf.csRGB)
    lines.append("render OK: %dx%d" % (pix.width, pix.height))
    t = d[2].get_text()
    lines.append("text layer len = %d" % len(t))
    d.close()
except Exception:
    lines.append("pymupdf FAILED:\n" + traceback.format_exc())

# 2) openpyxl
try:
    import openpyxl
    wb = openpyxl.load_workbook(r"D:\BaiduNetdiskDownload\北宋杨家将\动画改编方案\杨家将_分镜头脚本表.xlsx")
    lines.append("openpyxl OK: " + openpyxl.__version__ + " | sheets = " + " / ".join(wb.sheetnames))
    ws = wb["序幕分镜48镜"]
    lines.append("sheet rows = %d" % ws.max_row)
except Exception:
    lines.append("openpyxl FAILED:\n" + traceback.format_exc())

# 3) pypdf
try:
    import pypdf
    lines.append("pypdf OK: " + pypdf.__version__)
except Exception:
    lines.append("pypdf FAILED:\n" + traceback.format_exc())

# 4) site-packages 完好性：统计 pymupdf 目录文件数
sp = r"C:\Users\ken2zhao\.workbuddy\binaries\python\envs\default\Lib\site-packages\pymupdf"
n = 0
for root, dirs, files in os.walk(sp):
    n += len(files)
lines.append("pymupdf dir file count = %d" % n)

with open(LOG, "w", encoding="utf-8") as fh:
    fh.write("\n".join(lines))
