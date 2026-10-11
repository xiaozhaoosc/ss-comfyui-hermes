# -*- coding: utf-8 -*-
import os, re, traceback
BASE = r"D:\BaiduNetdiskDownload\北宋杨家将\动画改编方案"
OUT = r"D:\BaiduNetdiskDownload\北宋杨家将\.workbuddy\_tmp\verify2.log"
L = []
html = os.path.join(BASE, "杨家将_剧本与分镜全本.html")
s = open(html, "r", encoding="utf-8").read()
L.append("HTML size = %.1f KB" % (len(s.encode("utf-8"))/1024))
for tag in ("section", "main", "footer", "table", "tbody", "thead"):
    L.append("  <%s> open=%d  close=%d" % (tag, len(re.findall(r"<%s[ >]" % tag, s)), len(re.findall(r"</%s>" % tag, s))))
L.append("  div open=%d close=%d" % (len(re.findall(r"<div[ >]", s)), len(re.findall(r"</div>", s))))
L.append("  集块 .ep = %d" % len(re.findall(r'<div class="ep">', s)))
L.append("  场景块 .sc = %d" % len(re.findall(r'<div class="sc">', s)))
L.append("  镜头 △ = %d" % s.count('<p class="act">'))
L.append("  单元格 <td = %d" % len(re.findall(r"<td[ >]", s)))
L.append("  页标记 pg = %d" % len(re.findall(r'class="pg"', s)))
L.append("  残留 APPEND = %d" % s.count("APPEND"))
# 页标记覆盖的册-页
pgs = set(re.findall(r'<span class="pg">(\d-\d+)</span>', s))
L.append("  distinct 册-页 标记 = %d" % len(pgs))
try:
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(BASE, "杨家将_逐页分镜总表.xlsx"))
    L.append("XLSX sheets = " + " / ".join(wb.sheetnames))
    for n in wb.sheetnames:
        L.append("  [%s] rows=%d cols=%d" % (n, wb[n].max_row, wb[n].max_column))
    ws = wb["逐页分镜总表672页"]
    filled = sum(1 for r in range(2, ws.max_row+1) if (ws.cell(row=r, column=8).value or "").strip())
    L.append("  有文字的行 = %d / %d" % (filled, ws.max_row-1))
except Exception:
    L.append("XLSX FAIL\n" + traceback.format_exc())
open(OUT, "w", encoding="utf-8").write("\n".join(L))
